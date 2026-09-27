import csv,gzip,json,collections,runpy,html
from pathlib import Path
from datetime import datetime,timezone,timedelta
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
ROOT=Path(__file__).resolve().parents[1];D=ROOT/'reports/sojam-20260914/intent-audit'
classify=runpy.run_path(str(ROOT/'tools/sojam-intent-audit-build.py'))['classify']
def load(n):return json.loads((D/n).read_text(encoding='utf8'))
def save(n,o):(D/n).write_text(json.dumps(o,ensure_ascii=False,indent=2),encoding='utf8')
def kst(s):return datetime.fromisoformat(s.replace('Z','+00:00')).astimezone(timezone(timedelta(hours=9))).strftime('%m/%d %H:%M:%S')
def rng(v):
    v=sorted(set(x for x in v if x and x>0));return None if not v else str(v[0])if len(v)==1 else str(v[0])+'~'+str(v[-1])
u=load('universe.json');inventory={x['id']:x for x in load('inventory.json')};stats=load('today_stats.json');live=load('state_live.json')
stat={x['id']:x for b in stats['results']for x in b['response']['data']};requested={i for b in stats['results']for i in b['ids']}
assert not stats['failed']
assert set(load('candidate_ids.json'))<=requested
assert not classify('머리두피가려움')['pain'] and not classify('한포진물집')['pain']
assert classify('몸이너무가려워요')['pain']and classify('항문소양증치료')['priority']
assert classify('두드러기치료')['category']=='기존 운영제한'
assert classify('백반증치료')['priority']
groups={r[1]:r for r in csv.reader((D/'master_Adgroup.tsv').open(encoding='utf8'),delimiter='\t')if r}
ads=collections.defaultdict(list)
for r in csv.reader((D/'master_Ad.tsv').open(encoding='utf8'),delimiter='\t'):
    if r:ads[r[1]].append(r)
camps={c['nccCampaignId']:c for c in load('campaigns.json')}
glive={g['nccAdgroupId']:g for g in live['groups']};alive={g['gid']:g['rows']for g in live['ads']};klive={k['nccKeywordId']:k for k in load('keywords_live.json')}
discovery=load('discovery.tsv');vol={k['relKeyword']:k for k in discovery['keywords']}
serps={r['term']:r for r in load('serp/observations.json')if 'topAdCount'in r}
def state(kid):
    k=inventory[kid];g=groups[k['gid']];c=camps.get(g[2],{});lk=klive.get(kid);lg=glive.get(k['gid'])
    if lk and lk.get('delFlag'):return '등록 삭제'
    if (lk['userLock']if lk else k['lock']):return '키워드 OFF'
    if c.get('userLock'):return '캠페인 OFF'
    if (lg['userLock']if lg else g[5]=='1'):return '그룹 OFF'
    if lk and lk['status']!='ELIGIBLE':return lk['statusReason']
    if k['gid']in alive:
        if not any(not a['userLock']and not a.get('delFlag')and a['status']=='ELIGIBLE'and a['inspectStatus']=='APPROVED'for a in alive[k['gid']]):return '승인·노출가능 소재 없음'
        if lk:return '노출가능 상태 확인'
    if not ads[k['gid']]:return '소재 등록 없음'
    if not any(a[8]=='0'for a in ads[k['gid']]):return '소재 전부 OFF'
    return 'ON 구조 확인·소재 승인 추가확인'
rows=[]
for r in u:
    a=[stat[i]for i in r['ids']if i in stat];im=sum(x['impCnt']for x in a);ss=collections.Counter(state(i)for i in r['ids']);sv=serps.get(r['keyword']);v=vol.get(r['keyword'],{})
    todayrank=round(sum(x['impCnt']*x['avgRnk']for x in a)/im,2)if im else None
    if not r['keywordOn']:issue='등록 전부 OFF'
    elif not r['week']['imp']:issue='최근7일 노출 0·수요와 차단원인 점검'
    elif not im:issue='오늘 집계 노출 없음'
    elif todayrank>5:issue='오늘 평균순위 5위 밖'
    elif im<10:issue='오늘 노출 10회 미만·표본 적음'
    else:issue='오늘 노출 확인'
    if not sv:obs='미측정'
    elif sv.get('ageFilter'):obs='연령확인 필터·판정불가'
    elif sv['rank']:obs=f"상단 광고 {sv['rank']}위"
    elif sv['topAdCount']:obs=f"상단 {sv['topAdCount']}개 광고에서 미관찰"
    else:obs='상단 광고영역 미관찰'
    out={'키워드':r['keyword'],'의도분류':r['category'],'질환축':r['axis'],'간절함근거':r['signals'],'판정설명':r['reason'],'우선대상':r['priority'],'지역':r['region'],'등록수':r['registered'],'키워드ON수':r['keywordOn'],'구조진단':' / '.join(f'{k} {v}등록'for k,v in ss.items()),'어제노출':r['yesterday']['imp'],'어제클릭':r['yesterday']['clicks'],'어제비용':r['yesterday']['cost'],'어제전체매체평균순위':r['yesterday']['rank'],'최근7일노출':r['week']['imp'],'최근7일클릭':r['week']['clicks'],'최근7일비용':r['week']['cost'],'최근7일전체매체평균순위':r['week']['rank'],'오늘조회대상':all(i in requested for i in r['ids']),'오늘노출':im if all(i in requested for i in r['ids'])else None,'오늘클릭':sum(x['clkCnt']for x in a)if all(i in requested for i in r['ids'])else None,'오늘전체매체평균순위':todayrank,'오늘모바일통검순위범위':rng([x.get('mblNxAvgRnk')for x in a]),'오늘PC통검순위범위':rng([x.get('pcNxAvgRnk')for x in a]),'모바일직접관찰':obs,'직접관찰시각':kst(sv['observedAt'])if sv else None,'직접관찰위치':'강남 위치 미검증'if sv else None,'월PC검색수':v.get('monthlyPcQcCnt'),'월모바일검색수':v.get('monthlyMobileQcCnt'),'점검사항':issue,'키워드ID':' / '.join(r['ids'])}
    rows.append(out)
priority=[r for r in rows if r['우선대상']];clicked=sorted([r for r in rows if r['어제클릭']],key=lambda r:-r['어제비용'])
def order(r):return(-int('오늘 평균순위 5위 밖'in r['점검사항']),-r['어제클릭'],-r['최근7일노출'])
priority.sort(key=order)
new=[];registered={r['키워드']for r in rows}
for d in discovery['keywords']:
    c=classify(d['relKeyword'])
    if c['priority']:
        new.append({'키워드':d['relKeyword'],'의도분류':c['category'],'질환축':c['axis'],'등록여부':d['relKeyword']in registered,'월PC검색수':d['monthlyPcQcCnt'],'월모바일검색수':d['monthlyMobileQcCnt'],'시드':' / '.join(d['seeds'])})
new.sort(key=lambda r:r['등록여부'])
counts=collections.Counter(r['의도분류']for r in priority)
summary={'sourceDate':'2026-09-14','registered':len(inventory),'unique':len(rows),'priorityUnique':len(priority),'priorityCategories':dict(counts),'sevenDayNoImpression':sum(not r['최근7일노출']for r in priority),'sevenDayExposed':sum(bool(r['최근7일노출'])for r in priority),'todayExposed':sum(bool(r['오늘노출'])for r in priority),'allKeywordOff':sum(not r['키워드ON수']for r in priority),'todayRankOver5':sum((r['오늘전체매체평균순위']or 0)>5 for r in priority),'todayRequestedIds':len(requested),'todayReturnedIds':len(stat),'statsFetchedFrom':min(b['fetchedAt']for b in stats['results']),'statsFetchedUntil':max(b['fetchedAt']for b in stats['results']),'statsCompTm':sorted({b['response']['compTm']for b in stats['results']}),'statsCycleBaseTm':sorted({b['response']['cycleBaseTm']for b in stats['results']}),'directlyObservedKeywords':len(serps),'directAdsFound':sum(bool(r.get('rank'))for r in serps.values()),'discoveredKeywords':len(discovery['keywords']),'discoverySeeds':len(discovery['seeds']),'discoveryErrors':discovery['errors'],'newPriority':sum(not r['등록여부']for r in new),'stateKeywords':len(klive),'stateGroups':len(glive),'stateErrors':live['failed'],'yesterdayClicks':sum(r['어제클릭']for r in clicked),'yesterdayCost':sum(r['어제비용']for r in clicked),'priorityClickedWords':sum(r['우선대상']for r in clicked),'priorityClicks':sum(r['어제클릭']for r in clicked if r['우선대상']),'priorityCost':sum(r['어제비용']for r in clicked if r['우선대상'])}
save('summary.json',summary);save('review_rows.json',priority);save('new_candidates.json',new)
def csvwrite(file,data,compressed=False):
    if not data:return
    op=gzip.open if compressed else open
    with op(D/file,'wt',encoding='utf-8-sig',newline='')as f:
        w=csv.DictWriter(f,fieldnames=list(data[0]));w.writeheader();w.writerows(data)
csvwrite('전체63858검색어_의도노출.csv.gz',rows,True);csvwrite('우선키워드_노출순위.csv',priority);csvwrite('어제클릭58검색어.csv',clicked)
notes=[
 {'항목':'조사 범위','설명':f'전체 {len(inventory):,}등록 / {len(rows):,}고유어를 규칙 전수 선별. 우선 후보 {len(priority):,}고유어. 전부를 사람이 개별 정밀 판독했다는 뜻은 아님.'},
 {'항목':'실제 내원','설명':'최신 상담 시트 13탭 478문의행 / 정확한 내원 표시 195행. 키워드ID 연결이 없어 키워드별 실제 내원·결제 전환율은 산출하지 않음. 질환별 표는 중복 질환 포함.'},
 {'항목':'실시간과 평균순위','설명':'직접 검색은 9/14 13:25~13:26 모바일 비로그인 16개. 브라우저 위치는 역삼으로 요청했지만 검색 위치 확인 실패. 동일 시각 모든 이용자의 순위나 강남 확정순위가 아님.'},
 {'항목':'오늘 집계 시각','설명':'오늘 통계 13:01~13:02 조회, cycleBaseTm=12:00, compTm=12:43. recentAvgRnk는 전부 0으로 반환되어 유효한 실시간 순위로 사용하지 않음.'},
 {'항목':'매체 구분','설명':'전체매체 평균순위와 모바일/PC 통합검색 평균순위를 구분. 모바일/PC 통검 범위는 중복 등록별 유효한 값의 최소~최대이며 가중 평균이 아님. 0은 0위가 아닌 유효순위 없음.'},
 {'항목':'노출 0 해석','설명':'7일 노출 0은 실제 보고서 값. 검색량 부족·등록시점·지역·시간·매체·입찰·중지·소재 문제를 구분해야 하며 자동으로 입찰 부족이라 결론내리지 않음.'},
 {'항목':'소재 확인 범위','설명':f'{len(glive)}그룹/{len(klive)}키워드는 상태·소재 승인·타기팅을 API 재확인. 나머지는 전체 마스터의 ON/OFF와 소재 등록 구조를 확인했으며 승인·시간·지역 전체 검증 완료로 표시하지 않음.'},
 {'항목':'추가 발굴','설명':f'최신 키워드 도구 {len(discovery["seeds"])}시드에서 {len(discovery["keywords"]):,}연관어를 수집. 후보 {len(new)}개 중 미등록 {summary["newPriority"]}개. 모든 가능한 자연어 표현을 완전히 발굴했다는 뜻은 아님. <10은 0이 아님.'},
 {'항목':'과거 진료범위 제한','설명':'두드러기·건선 등에 실제 내원 이력이 있어도 기존 제외 지시는 보존. 백반증·다한증은 유지 축으로 포함. 광고 설정을 변경하지 않음.'},
 {'항목':'어제 대조','설명':'AD_DETAIL 41,440노출/79클릭/127,990원은 캠페인 합계와 정확히 일치. 등록어 58개는 78클릭/124,335원, 플레이스 1클릭/3,655원. EXPKEYWORD 78클릭/124,334원으로 반올림 1원 차이.'},
 {'항목':'네이버 공식 순위 정의','설명':'https://github.com/naver/searchad-apidoc/wiki/FAQ-stat'},
]
wb=Workbook();wb.remove(wb.active)
def sheet(name,data):
    ws=wb.create_sheet(name)
    if not data:return
    keys=list(data[0]);ws.append(keys)
    for r in data:ws.append([str(r.get(k))if isinstance(r.get(k),(list,dict))else r.get(k)for k in keys])
    ws.freeze_panes='B2';ws.auto_filter.ref=ws.dimensions
    for c in ws[1]:c.fill=PatternFill('solid',fgColor='173E43');c.font=Font(color='FFFFFF',bold=True)
    for i,k in enumerate(keys,1):ws.column_dimensions[ws.cell(1,i).column_letter].width=50 if k in ['설명','구조진단','판정설명','키워드ID']else 28 if i==1 else 21
    if name=='읽는법':
        ws.column_dimensions['B'].width=110
        for row in ws.iter_rows(min_row=2):row[1].alignment=Alignment(wrap_text=True,vertical='top');ws.row_dimensions[row[0].row].height=44
sheet('읽는법',notes);sheet('어제 클릭58검색어',clicked);sheet('우선후보4539',priority);sheet('7일 노출없음',[r for r in priority if not r['최근7일노출']]);sheet('전체등록OFF',[r for r in priority if not r['키워드ON수']]);sheet('오늘 평균5위밖',[r for r in priority if(r['오늘전체매체평균순위']or 0)>5]);sheet('추가발굴',new)
sheet('직접 모바일 관찰',[{'키워드':r['term'],'시각':kst(r['observedAt']),'소잠상단광고순위':r['rank'],'상단광고수':r['topAdCount'],'연령필터':r['ageFilter'],'위치검증':'불가','해석':'특정 측정환경의 순간 관찰; 미관찰은 전체 미노출이 아님'}for r in serps.values()])
sheet('실제 내원 질환별',list(csv.DictReader((ROOT/'reports/sojam-20260908/sheet-linked/diseases.csv').open(encoding='utf-8-sig'))))
sheet('운영제한 클릭',[r for r in clicked if r['의도분류']in ['기존 운영제한','진료적합성 확인']])
book=D/'소잠_내원의도_노출순위_20260914.xlsx';wb.save(book)
check=load_workbook(book,read_only=True);assert check['어제 클릭58검색어'].max_row==59;assert check['우선후보4539'].max_row==len(priority)+1;check.close()
md=['# 소잠 내원 의도·노출·순위 점검 — 2026-09-14','',f'전체 {len(inventory):,}등록/{len(rows):,}고유 검색어에서 우선 후보 {len(priority):,}개를 선별했습니다. 실제 내원 확률을 산출한 자료가 아닙니다.','',f'어제 클릭은 등록어 58개 78클릭 + 플레이스 1클릭입니다. 명시적인 치료·내원 탐색 또는 강한 고통 표현은 {summary["priorityClickedWords"]}개 검색어 / {summary["priorityClicks"]}클릭 / {summary["priorityCost"]:,.0f}원입니다.','', '|어제 우선 검색어|클릭|비용|직접 모바일 관찰|','|---|---:|---:|---|']
for r in clicked:
    if r['우선대상']:md.append(f'|{r["키워드"]}|{r["어제클릭"]}|{r["어제비용"]:,.0f}원|{r["모바일직접관찰"]}|')
md+=['',f'최근 7일(9/7~13) 노출 있음 {summary["sevenDayExposed"]:,}개, 노출 없음 {summary["sevenDayNoImpression"]:,}개. 등록 전부 OFF {summary["allKeywordOff"]}개는 노출 없는 후보와 겹칠 수 있습니다. 오늘 집계 노출 있음 {summary["todayExposed"]}개, 오늘 전체매체 평균순위 5위 밖 {summary["todayRankOver5"]}개입니다.','', '### 우선 점검할 순위','', '|키워드|최근7일 노출|오늘 노출|오늘 전체매체 평균순위|오늘 모바일 통검 순위|','|---|---:|---:|---:|---|']
for r in priority:
    if(r['오늘전체매체평균순위']or 0)>5:md.append(f'|{r["키워드"]}|{r["최근7일노출"]}|{r["오늘노출"]}|{r["오늘전체매체평균순위"]}|{r["오늘모바일통검순위범위"]or"유효순위 없음"}|')
md+=['','### 해석과 검증','']+[f'- **{n["항목"]}**: {n["설명"]}'for n in notes]
md+=['','[전체 결과 엑셀](소잠_내원의도_노출순위_20260914.xlsx) · [우선 후보 CSV](우선키워드_노출순위.csv) · [전체 검색어 CSV 압축](전체63858검색어_의도노출.csv.gz)','']
(D/'결과.md').write_text('\n'.join(md),encoding='utf8')
print(json.dumps(summary,ensure_ascii=False))
