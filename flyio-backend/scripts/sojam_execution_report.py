import csv,json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill
ROOT=Path(__file__).resolve().parents[2];OUT=ROOT/'reports/sojam-20260908/execution'
def read(n):return json.loads((OUT/(n+'.json')).read_text(encoding='utf-8'))
def csvrows(n):return list(csv.DictReader((OUT/n).open(encoding='utf-8-sig')))
stats={r['id']:r for s in read('stats_device') for r in s.get('data',[])}
rows=csvrows('live_keywords.csv')
for r in rows:
    s=stats.get(r['keyword_id']);r['stats_returned']=s is not None
    for device in ['all','PC','모바일']:
        d=s if device=='all' else next((b for b in (s or {}).get('breakdowns',[]) if b.get('name')==device),None)
        for field in ['impCnt','clkCnt','salesAmt','avgRnk']:r[device+'_'+field]=d.get(field) if d else None
        r[device+'_actual_CPC']=round(d['salesAmt']/d['clkCnt'],2) if d and d.get('clkCnt') else None
with (OUT/'live_keywords_with_device_stats.csv').open('w',encoding='utf-8-sig',newline='') as f:
    w=csv.DictWriter(f,fieldnames=list(rows[0]));w.writeheader();w.writerows(rows)
selected=set(r['keyword_id'] for r in read('shopping_pause_candidates'))
summary=dict(manifest=read('manifest'),period='2026-09-01 ~ 2026-09-07',stats_keyword_rows=len(stats),devices={},shopping_historical_spend=sum(r.get('salesAmt',0) for k,r in stats.items() if k in selected),shopping_historical_clicks=sum(r.get('clkCnt',0) for k,r in stats.items() if k in selected))
for dev in ['PC','모바일']:
    ds=[b for s in stats.values() for b in s.get('breakdowns',[]) if b.get('name')==dev]
    v={f:sum(x.get(f,0) for x in ds) for f in ['impCnt','clkCnt','salesAmt']};v['actual_CPC']=round(v['salesAmt']/v['clkCnt'],2) if v['clkCnt'] else None;summary['devices'][dev]=v
cs=read('campaigns')['items'];summary['daily_budget_total']=sum(c.get('dailyBudget',0) for c in cs if c.get('useDailyBudget'))
summary['campaign_budget_cap_below_group_budget']=[]
for g in read('selected_groups'):
    c=next(c for c in cs if c['nccCampaignId']==g['nccCampaignId'])
    if c.get('useDailyBudget') and g.get('useDailyBudget') and g['dailyBudget']>c['dailyBudget']:
        summary['campaign_budget_cap_below_group_budget'].append(dict(group=g['name'],campaign=c['name'],groupBudget=g['dailyBudget'],campaignBudget=c['dailyBudget']))
(OUT/'analysis_summary.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')
wb=Workbook();wb.remove(wb.active)
campaign_rows=[{k:c.get(k) for k in ['nccCampaignId','name','campaignTp','status','userLock','useDailyBudget','dailyBudget','budgetLock']} for c in cs]
tables={'검증범위':[{'항목':k,'값':str(v)} for k,v in summary['manifest'].items()],'캠페인예산':campaign_rows,'현재입찰_기기별실적':rows,'상품검색어중지':read('shopping_pause_candidates'),'소재노출조건':csvrows('ad_delivery_groups.csv')}
if (OUT/'shopping_pause_result.json').exists(): tables['변경후검증']=read('shopping_pause_result')['verified']
for title,rs in tables.items():
    ws=wb.create_sheet(title)
    if not rs:continue
    keys=list(rs[0]);ws.append(keys)
    for r in rs:ws.append([r.get(k) for k in keys])
    ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
    for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='244062')
    for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=min(48,max(15,len(str(col[0].value))+3))
wb.save(OUT/'실행검증_입찰예산.xlsx')
print(json.dumps(summary,ensure_ascii=False));print('TOP PC')
for r in sorted(rows,key=lambda r:r['PC_salesAmt'] or 0,reverse=True)[:15]:print(r['keyword'],r['pc_effective_bid'],r['PC_salesAmt'],r['PC_clkCnt'],r['PC_actual_CPC'],r['PC_avgRnk'])
