import csv,json,gzip,tempfile,os,shutil,tarfile
from pathlib import Path
from collections import Counter
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font,PatternFill
ROOT=Path(__file__).resolve().parents[2];P=ROOT/'reports/sojam-20260909';D=P/'plan';RESULT=P/'verified/sojam-apply-20260909'
def load(p):return json.loads(p.read_text(encoding='utf-8'))
def main():
 s=load(D/'summary.json');plan=load(D/'plan.json')
 with tarfile.open(P/'verified.tar.gz','r:gz') as archive:result=json.load(archive.extractfile('sojam-apply-20260909/result.json'))
 assert result.get('complete') and len(result['keywords'])==len(plan['keywords']) and all(r['verified'] for r in result['keywords'])
 assert len(result['budgets'])==len(plan['budgets']) and all(r['verified'] for r in result['budgets'])
 assert result['daily_budget_total_after']==156750
 counts=Counter(r['action'] for r in result['keywords']);wb=Workbook(write_only=True)
 def table(title,keys,rows):
  ws=wb.create_sheet(title);ws.freeze_panes='A2';head=[]
  from openpyxl.utils import get_column_letter
  for i,k in enumerate(keys,1):ws.column_dimensions[get_column_letter(i)].width=65 if k=='reason' else 38 if k in ['keyword','group_name','campaign_name','name'] else 20
  for k in keys:
   c=WriteOnlyCell(ws,k);c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='244062');head.append(c)
  ws.append(head);n=1
  for r in rows:
   values=[]
   for k in keys:
    v=r.get(k)
    if isinstance(v,str) and v and any(x in k for x in ['bid','weight','cap','cost','clicks','impressions','score','Budget','budget','rank','estimate']):
     try:v=float(v)
     except ValueError:pass
    values.append(v)
   ws.append(values);n+=1
  ws.auto_filter.ref=f'A1:{get_column_letter(len(keys))}{n}'
 table('검증요약',['항목','값'],[{'항목':'전체 등록행','값':s['rows']},{'항목':'고유 검색어','값':s['unique_keywords']},{'항목':'수집 그룹','값':s['source']['groups']},{'항목':'검증된 변경','값':len(result['keywords'])},{'항목':'유지 판정','값':s['actions']['KEEP']},{'항목':'일예산 합계','값':156750},{'항목':'결제 귀속','값':'미확인'},{'항목':'적용 완료 UTC','값':result['finished']}])
 for title,name in [('전체등록행_판정','전체키워드_조정판정.csv'),('입찰중지_변경안','실제변경대상.csv'),('캠페인예산_전체','캠페인예산_성과대조.csv'),('예산변경','캠페인예산_변경안.csv')]:
  reader=gzip.open(D/(name+'.gz'),'rt',encoding='utf-8-sig') if (D/(name+'.gz')).exists() else (D/name).open(encoding='utf-8-sig')
  with reader as f:
   rs=csv.DictReader(f);table(title,rs.fieldnames,rs)
 table('키워드변경후검증',list(result['keywords'][0]),result['keywords']);table('예산변경후검증',list(result['budgets'][0]),result['budgets'])
 fd,tmpname=tempfile.mkstemp(prefix='sojam-report-',suffix='.xlsx');os.close(fd)
 try:
  wb.save(tmpname);shutil.copyfile(tmpname,P/'전체키워드_입찰예산_적용검증.xlsx')
 finally:Path(tmpname).unlink(missing_ok=True)
 examples=[]
 for k in ['겨드랑이냄새','피부가려움증','지루성피부염한의원','아토피치료','주부습진치료']:
  candidates=[r for r in plan['keywords'] if r['keyword']==k]
  for r in sorted(candidates,key=lambda r:-r['period_cost'])[:2]:examples.append(f"| {k} | {r['group_name']} | {r['grade']} | {r['before_pc_bid']:,.0f} → {r['after_pc_bid']:,.0f}원 | {r['before_mobile_bid']:,.0f} → {r['after_mobile_bid']:,.0f}원 |")
 budgetlines=[f"| {b['name']} | {b['before']:,}원 | {b['after']:,}원 |" for b in plan['budgets']]
 report=f'''# 소잠 전체 키워드 입찰·예산 적용 결과

적용 완료 UTC: {result['finished']}. 고객번호 1858907.

전체 137개 캠페인·{s['source']['groups']:,}개 그룹과 키워드 마스터 {s['rows']:,}행(고유 검색어 {s['unique_keywords']:,}개)을 확인했다. 전수 정책 판정과 지출 상위·중지 후보 개별 검토를 거쳤다. 모든 고유 검색어를 사람이 개별 정밀 검토했다는 뜻은 아니다. 명확하지 않은 검색어는 U로 남겼으며, 결제 가능성을 추정하지 않았다.

## 실제 적용 및 검증

- 입찰 인하 {counts['BID_DOWN']:,}개, 입찰 인상 {counts['BID_UP']:,}개, 추가 중지 {counts['PAUSE']:,}개. 등록 ID 기준이다.
- 나머지 {s['actions']['KEEP']:,}행은 유지 판정. 기존 OFF·상위 OFF·타 광고 유형·브랜드·상한 이내 등 사유가 전수표에 있다. 9/8 중지한 72개를 이번 신규 중지 수에 더하지 않는다.
- 캠페인 예산 7개 변경. 계정 일예산 사용 설정 합계 156,750원 유지. 17,000원을 기존 캠페인 사이에서 재배분했다. 브랜드검색 계약 비용은 별도다.
- 변경 직전 실제 키워드·그룹·계정·입찰·가중치·상속 여부를 대조했다. 모든 변경은 별도 GET으로 결과를 확인했다. OFF 항목을 ON으로 바꾼 작업은 없다.
- 증액 그룹은 승인된 ELIGIBLE 소재를 확인했다. 의료광고 소재·랜딩·타기팅은 이번에 변경하지 않았다.

## 입찰 기준

| 등급 | 기준 | 기기 가중치 반영 입찰 상한 |
|---|---|---:|
| S | 지속·재발·부담 표현과 치료 탐색 동시 확인 | 4,500원 |
| A | 피부질환 치료기관·진료 탐색 | 3,500원 |
| B | 지속 증상·생활 불편, 치료기관 선택 미확인 | 2,200원 |
| C | 핵심 질환·증상명 | 1,500원 |
| D | 핵심 질환의 원인·자가관리 정보 | 900원 |

다른 피부질환 치료 탐색은 2,000원, 기타 피부 증상은 1,000원, 진균성 정보는 700원으로 제한했다. 기존 확대 보류 범위·의미 미확정 검색어는 최대 500원, 전문·긴급 진료 적합성 우선 범위는 300원으로 제한했다. 이 숫자는 결제 데이터로 최적화된 정답이 아니라 총예산 안에서 검증할 초기 운영 상한이다. 원거리 검색어에는 별도 제한을 적용했다.

기본입찰에 PC·모바일 가중치를 곱해 상한을 확인했다. 이미 상한 이내면 기계적으로 올리지 않았다. 20회 이상 노출됐고 평균순위가 4보다 낮은 치료형은 약 10% 증액했다. 최근 무노출 치료형 중 네이버 28일 최소노출 추정치가 현재 입찰보다 높고 소액 시험 한도 안에 들어오는 항목은 그 추정치에 맞췄다(S 기본입찰 1,500원/A 1,000원 이내). 추정치 70원만으로 검색량이나 실제 노출을 보장하지 않는다. 기존 운영 프로필의 제외어는 이 무노출 증액에서 제외했다.

## 실제 변경 예시

아래 금액은 현재 설정 가중치를 반영한 계산입찰이며 실제 과금 CPC가 아니다. 같은 검색어라도 그룹이 다르면 서로 다른 등록이다.

| 키워드 | 그룹 | 등급 | PC 적용입찰 | 모바일 적용입찰 |
|---|---|---|---:|---:|
{chr(10).join(examples)}

## 예산 변경

| 캠페인 | 이전 일예산 | 변경 일예산 |
|---|---:|---:|
{chr(10).join(budgetlines)}

최근 치료형 검색어 지출이 확인되는 주력·대표 캠페인에 배분을 늘렸다. 캠페인 이름만으로 결정하지 않고 내부 키워드 등급과 9/1~8 지출을 대조했다. 다만 치료형 지출이 곧 내원이나 결제 성과라는 뜻은 아니며, 그룹 구성이 혼합돼 있어 캠페인 예산을 치료형 전용 예산이라고 부를 수 없다. 인하를 먼저 적용하고 인상을 나중에 적용했다.

## 성과 기준 및 한계

9/1~8 대량 보고서 총 236,060노출·430클릭·1,145,678원. 별도 캠페인 실적 API와 노출·클릭이 일치하고 비용은 1원 차이였다. 등록 키워드로 귀속되지 않은 25,606원/7클릭은 별도 보존했다.

시트 13탭 최신 확인: 문의 464행·내원 표시 191행. 실제 첫내원일·키워드ID·수납금액 연결은 아직 없다. 따라서 이번 변경이 월 50신환 또는 500만원 결제를 달성했다고 말할 수 없다. 별점·동의를 결제로 대체하지 않았다.

변경 직후 성과는 아직 없다. 현재 고비용 키워드 중 일부를 크게 인하했으므로 노출·문의가 줄 수 있다. 48~72시간에는 모바일 노출·클릭·예산 사용량과 문의를 먼저 점검하고, 7일에는 실제 첫내원과 문의→내원율을 대조한다. 치료형 지출도 내원이 발생하지 않으면 증액하지 않는다. 결제 원장이 연결되면 순결제 CPA까지 판단 기준을 확장한다. 절감액·ROAS·50내원을 미리 확정하지 않는다.

서버의 도메인 자동 발굴과 자동 입찰 설정은 조회 당시 모두 꺼져 있었다. 이번 정책의 상시 자동 적용이나 예약 재검증을 새로 설정한 것은 아니다. 소스 시트에 직접 쓰기·결제 연동도 미완료다.

API 기준: [네이버 공식 입찰·최소노출 추정 예제](https://github.com/naver/searchad-apidoc/blob/master/python-sample/examples/ad_management_sample.py), [공식 마스터 보고서 필드 정의](https://gist.github.com/naver-searchad/186ca42e1e8596b0e3dcf74e3a86c04f).
'''
 (P/'전체변경_결과.md').write_text(report,encoding='utf-8')
 print(json.dumps(dict(verified=len(result['keywords']),actions=dict(counts),budget=result['daily_budget_total_after']),ensure_ascii=False))
if __name__=='__main__':main()
